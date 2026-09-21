"""
FHIRSheets XLSX Toolkit for LangGraph Agent

This module provides a toolkit of tools for working with FHIR spreadsheets.
"""

import json
import logging
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Type

import openpyxl
from langchain_core.tools import BaseTool, BaseToolkit
from pydantic import BaseModel, Field

logger = logging.getLogger(__name__)


# Workbook Manager for shared instance (internal use only)
class WorkbookManager:
    """Manages a shared workbook instance across multiple tool calls."""
    
    _instance: Optional['WorkbookManager'] = None
    _workbook: Optional[openpyxl.Workbook] = None
    _file_path: Optional[str] = None
    
    def __new__(cls):
        if cls._instance is None:
            cls._instance = super(WorkbookManager, cls).__new__(cls)
        return cls._instance
    
    @classmethod
    def open_workbook(cls, file_path: str) -> Dict[str, Any]:
        """Open a workbook and store it in the shared instance."""
        manager = cls()
        
        # If a different workbook is already open, close it first
        if manager._workbook is not None and manager._file_path != file_path:
            logger.info(f"Closing currently open workbook: {manager._file_path}")
            cls.close_workbook(save=True)
        
        # If the same workbook is already open, just return success
        if manager._workbook is not None and manager._file_path == file_path:
            logger.info(f"Workbook already open: {file_path}")
            return {
                "success": True,
                "message": f"Workbook '{file_path}' is already open",
                "file_path": file_path
            }
        
        try:
            manager._workbook = openpyxl.load_workbook(file_path)
            manager._file_path = file_path
            sheet_names = manager._workbook.sheetnames
            logger.info(f"Opened workbook: {file_path} with sheets: {sheet_names}")
            return {
                "success": True,
                "message": f"Successfully opened workbook: {file_path}",
                "file_path": file_path,
                "sheet_names": sheet_names
            }
        except FileNotFoundError:
            return {
                "success": False,
                "error": f"File not found: {file_path}"
            }
        except Exception as e:
            logger.error(f"Error opening workbook: {e}")
            return {
                "success": False,
                "error": f"Error opening workbook: {str(e)}"
            }
    
    @classmethod
    def close_workbook(cls, save: bool = True) -> Dict[str, Any]:
        """Close the currently open workbook."""
        manager = cls()
        
        if manager._workbook is None:
            return {
                "success": False,
                "message": "No workbook is currently open"
            }
        
        try:
            file_path = manager._file_path
            if save:
                manager._workbook.save(file_path)
                logger.info(f"Saved and closed workbook: {file_path}")
                message = f"Successfully saved and closed workbook: {file_path}"
            else:
                logger.info(f"Closed workbook without saving: {file_path}")
                message = f"Closed workbook without saving: {file_path}"
            
            manager._workbook.close()
            manager._workbook = None
            manager._file_path = None
            
            return {
                "success": True,
                "message": message,
                "file_path": file_path
            }
        except Exception as e:
            logger.error(f"Error closing workbook: {e}")
            return {
                "success": False,
                "error": f"Error closing workbook: {str(e)}"
            }
    
    @classmethod
    def get_workbook(cls) -> Optional[openpyxl.Workbook]:
        """Get the currently open workbook."""
        manager = cls()
        return manager._workbook
    
    @classmethod
    def get_file_path(cls) -> Optional[str]:
        """Get the file path of the currently open workbook."""
        manager = cls()
        return manager._file_path
    
    @classmethod
    def is_open(cls) -> bool:
        """Check if a workbook is currently open."""
        manager = cls()
        return manager._workbook is not None


# Naming validation utilities
def _to_snake_case(name: str) -> str:
    """Convert a string to snake_case."""
    # Replace spaces and hyphens with underscores
    name = name.replace(' ', '_').replace('-', '_')
    # Insert underscore before uppercase letters (for camelCase/PascalCase)
    name = re.sub('(.)([A-Z][a-z]+)', r'\1_\2', name)
    name = re.sub('([a-z0-9])([A-Z])', r'\1_\2', name)
    # Convert to lowercase and remove multiple underscores
    name = re.sub('_+', '_', name.lower())
    # Remove leading/trailing underscores
    return name.strip('_')


def _to_pascal_case(name: str) -> str:
    """Convert a string to PascalCase."""
    # Remove special characters and split on spaces, underscores, hyphens
    name = re.sub(r'[^a-zA-Z0-9\s_-]', '', name)
    # Split on spaces, underscores, or hyphens
    words = re.split(r'[\s_-]+', name)
    
    # If only one word (no splitting occurred), just capitalize first letter
    if len(words) == 1 and words[0]:
        word = words[0]
        return word[0].upper() + word[1:] if len(word) > 1 else word.upper()
    
    # Otherwise capitalize first letter of each word and join
    return ''.join(word.capitalize() for word in words if word)


# Shared state for current file tracking
class ToolkitState(BaseModel):
    """Shared state across toolkit tools."""
    current_file: Optional[Path] = Field(default=None, description="Currently active FHIR Excel file")
    file_history: List[Path] = Field(default_factory=list, description="History of files accessed")
    working_dir: Path = Field(default_factory=Path.cwd, description="Working directory for file operations")
    config_path: Optional[str] = Field(default=None, description="Path to configuration file")


# Pydantic Input Models
class CreateNewFileInput(BaseModel):
    """Input for creating a new FHIR cohort file."""
    file_name: Optional[str] = Field(default=None, description="Optional custom name for the new file")
    output_dir: Optional[str] = Field(default=None, description="Optional directory to save the file")


class SetCurrentFileInput(BaseModel):
    """Input for setting the current file."""
    file_path: str = Field(..., description="Path to the FHIR Excel file to set as current")


class GetCurrentFileInput(BaseModel):
    """Input for getting current file info."""
    pass


class CreateResourceDefinitionInput(BaseModel):
    """Input for creating a resource definition."""
    entity_name: str = Field(..., description="The name or identifier of the entity")
    resource_type: str = Field(..., description="The FHIR resource type")
    profiles: Optional[str] = Field(default=None, description="Comma-separated list of FHIR profile URLs")


class GetResourceDefinitionInput(BaseModel):
    """Input for getting a resource definition."""
    entity_name: str = Field(..., description="The name or identifier of the entity to retrieve")


class CreateResourceLinkInput(BaseModel):
    """Input for creating a resource link."""
    origin_resource: str = Field(..., description="The entity name that initiates the reference")
    reference_path: str = Field(..., description="The JSON path where the reference is made")
    destination_resource: str = Field(..., description="The entity name being referenced")


class AddPatientDataInput(BaseModel):
    """Input for adding patient data column."""
    entity_name: str = Field(..., description="The entity name")
    json_path: str = Field(..., description="The JSON path within the resource")
    data_type: str = Field(..., description="The FHIR data type")
    data_element: str = Field(..., description="Human-readable name for this data element")
    value_set: Optional[str] = Field(default=None, description="Optional ValueSet URL")
    recommended_profile: Optional[str] = Field(default=None, description="Optional recommended profile URL")


class CopyReferenceColumnsInput(BaseModel):
    """Input for copying reference columns."""
    entity_name: str = Field(..., description="The entity name from ResourceDefinitions")
    value_type: Optional[str] = Field(default=None, description="Optional value type filter")


class DeleteReferenceColumnsInput(BaseModel):
    """Input for deleting reference columns."""
    entity_name: str = Field(..., description="The entity name whose columns should be deleted")


class CheckResourceDefinitionExistsInput(BaseModel):
    """Input for checking if a resource definition exists."""
    entity_name: str = Field(..., description="The entity name to check")


class SetPatientDataValueInput(BaseModel):
    """Input for setting patient data value."""
    entity_name: str = Field(..., description="The entity name")
    data_element: str = Field(..., description="The data element name")
    row_index: int = Field(..., description="The 0-based patient data row index")
    value: str = Field(..., description="The value to set")


class GenerateFhirBundlesInput(BaseModel):
    """Input for generating FHIR bundles."""
    output_dir: str = Field(default="./output", description="Directory where output files will be saved")
    config_path: Optional[str] = Field(default=None, description="Optional path to configuration file")


class GetMetadataInput(BaseModel):
    """Input for getting metadata."""
    pass


# Tool Classes
class CreateNewFileTool(BaseTool):
    """Tool for creating a new FHIR cohort file from template."""
    
    name: str = "create_new_file"
    description: str = "Create a new FHIR cohort Excel file from the template and open it for editing"
    args_schema: Type[BaseModel] = CreateNewFileInput
    state: ToolkitState = Field(default_factory=ToolkitState)
    
    def _run(self, file_name: Optional[str] = None, output_dir: Optional[str] = None) -> str:
        try:
            target_dir = Path(output_dir) if output_dir else self.state.working_dir
            target_dir.mkdir(parents=True, exist_ok=True)
            
            if not file_name:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                file_name = f"fhir_cohort_{timestamp}.xlsx"
            else:
                # Convert file name to snake_case
                base_name = file_name.replace('.xlsx', '')
                snake_name = _to_snake_case(base_name)
                if snake_name != base_name:
                    logger.info(f"File name converted from '{base_name}' to '{snake_name}' (snake_case)")
                file_name = snake_name
            
            if not file_name.endswith('.xlsx'):
                file_name += '.xlsx'
            
            target_path = target_dir / file_name
            
            template_path = Path("src/resources/Fhir_Cohort_Import_Template.xlsx")
            if not template_path.exists():
                template_path = Path(__file__).parent.parent.parent / "resources" / "Fhir_Cohort_Import_Template.xlsx"
            
            if not template_path.exists():
                return json.dumps({"success": False, "error": "Template file not found"})
            
            shutil.copy2(template_path, target_path)
            self.state.current_file = target_path
            self.state.file_history.append(target_path)
            
            # Open the newly created file in WorkbookManager
            open_result = WorkbookManager.open_workbook(str(target_path))
            if not open_result.get("success"):
                return json.dumps({
                    "success": False,
                    "error": f"File created but failed to open: {open_result.get('error')}"
                })
            
            return json.dumps({
                "success": True,
                "file_path": str(target_path),
                "message": "Created new FHIR cohort file from template and opened it for editing"
            })
        except Exception as e:
            return json.dumps({"success": False, "error": str(e)})


class SetCurrentFileTool(BaseTool):
    """Tool for setting the current working file."""
    
    name: str = "set_current_file"
    description: str = "Set the current working file pointer to an existing FHIR Excel file and open it for editing"
    args_schema: Type[BaseModel] = SetCurrentFileInput
    state: ToolkitState = Field(default_factory=ToolkitState)
    
    def _run(self, file_path: str) -> str:
        try:
            path = Path(file_path)
            if not path.exists():
                return json.dumps({"success": False, "error": f"File not found: {file_path}"})
            if not path.suffix == '.xlsx':
                return json.dumps({"success": False, "error": "File must be an .xlsx file"})
            
            self.state.current_file = path
            if path not in self.state.file_history:
                self.state.file_history.append(path)
            
            # Open the file in WorkbookManager
            open_result = WorkbookManager.open_workbook(str(path))
            if not open_result.get("success"):
                return json.dumps({
                    "success": False,
                    "error": f"Failed to open file: {open_result.get('error')}"
                })
            
            return json.dumps({
                "success": True,
                "current_file": str(self.state.current_file),
                "message": "File set as current and opened for editing"
            })
        except Exception as e:
            return json.dumps({"success": False, "error": str(e)})


class GetCurrentFileTool(BaseTool):
    """Tool for getting current file information."""
    
    name: str = "get_current_file"
    description: str = "Get information about the current working file"
    args_schema: Type[BaseModel] = GetCurrentFileInput
    state: ToolkitState = Field(default_factory=ToolkitState)
    
    def _run(self) -> str:
        if self.state.current_file is None:
            return json.dumps({"has_current_file": False, "message": "No current file set"})
        
        return json.dumps({
            "success": True,
            "current_file": str(self.state.current_file),
            "file_name": self.state.current_file.name,
            "file_exists": self.state.current_file.exists(),
            "is_open_in_manager": WorkbookManager.is_open() and WorkbookManager.get_file_path() == str(self.state.current_file)
        })


class CreateResourceDefinitionTool(BaseTool):
    """Tool for creating resource definitions."""
    
    name: str = "create_resource_definition"
    description: str = "Create or update a resource definition entry in the ResourceDefinitions sheet"
    args_schema: Type[BaseModel] = CreateResourceDefinitionInput
    state: ToolkitState = Field(default_factory=ToolkitState)
    
    def _run(self, entity_name: str, resource_type: str, profiles: Optional[str] = None) -> str:
        try:
            # Convert entity name to PascalCase
            original_name = entity_name
            entity_name = _to_pascal_case(entity_name)
            if entity_name != original_name:
                logger.info(f"Entity name converted from '{original_name}' to '{entity_name}' (PascalCase)")
            
            # Use shared workbook instance
            wb = WorkbookManager.get_workbook()
            if wb is None:
                return json.dumps({"success": False, "error": "No workbook is currently open. Use set_current_file or create_new_file first."})
            
            if "ResourceDefinitions" not in wb.sheetnames:
                return json.dumps({"success": False, "error": "ResourceDefinitions sheet not found"})
            
            sheet = wb["ResourceDefinitions"]
            existing_row = None
            last_data_row = 2  # Start after header rows
            
            # Find existing entity and track last row with actual data
            for row in range(3, sheet.max_row + 1):
                cell_value = sheet.cell(row=row, column=1).value
                if cell_value == entity_name:
                    existing_row = row
                    break
                if cell_value is not None and str(cell_value).strip():
                    last_data_row = row
            
            target_row = existing_row if existing_row else last_data_row + 1
            action = "Updated" if existing_row else "Added"
            
            sheet.cell(row=target_row, column=1, value=entity_name)
            sheet.cell(row=target_row, column=2, value=resource_type)
            sheet.cell(row=target_row, column=3, value=profiles if profiles else "")
            
            # After creating/updating the resource definition, handle reference columns
            delete_result = {"success": True, "message": "Skipped - new resource"}
            copy_result = {"success": False, "message": "Skipped"}
            
            if action == "Added":
                # For new resources, just copy reference columns
                copy_tool = CopyReferenceColumnsTool(state=self.state)
                copy_result_str = copy_tool._run(entity_name=entity_name)
                copy_result = json.loads(copy_result_str)
            elif action == "Updated":
                # For updated resources, delete existing columns first
                delete_tool = DeleteReferenceColumnsTool(state=self.state)
                delete_result_str = delete_tool._run(entity_name=entity_name)
                delete_result = json.loads(delete_result_str)
                
                # Then copy fresh reference columns
                copy_tool = CopyReferenceColumnsTool(state=self.state)
                copy_result_str = copy_tool._run(entity_name=entity_name)
                copy_result = json.loads(copy_result_str)
            
            return json.dumps({
                "success": True,
                "entity_name": entity_name,
                "resource_type": resource_type,
                "action": action,
                "row": target_row,
                "delete_reference_columns": delete_result,
                "copy_reference_columns": copy_result
            })
        except Exception as e:
            return json.dumps({"success": False, "error": str(e)})


class GetResourceDefinitionTool(BaseTool):
    """Tool for getting resource definitions."""
    
    name: str = "get_resource_definition"
    description: str = "Get a resource definition entry from the ResourceDefinitions sheet"
    args_schema: Type[BaseModel] = GetResourceDefinitionInput
    state: ToolkitState = Field(default_factory=ToolkitState)
    
    def _run(self, entity_name: str) -> str:
        try:
            # Use shared workbook instance
            wb = WorkbookManager.get_workbook()
            if wb is None:
                return json.dumps({"success": False, "error": "No workbook is currently open. Use set_current_file or create_new_file first."})
            
            if "ResourceDefinitions" not in wb.sheetnames:
                return json.dumps({"success": False, "error": "ResourceDefinitions sheet not found"})
            
            sheet = wb["ResourceDefinitions"]
            
            # Find the entity
            for row in range(3, sheet.max_row + 1):
                cell_value = sheet.cell(row=row, column=1).value
                if cell_value == entity_name:
                    resource_type = sheet.cell(row=row, column=2).value
                    profiles = sheet.cell(row=row, column=3).value
                    
                    return json.dumps({
                        "success": True,
                        "entity_name": entity_name,
                        "resource_type": resource_type,
                        "profiles": profiles if profiles else "",
                        "row": row
                    })
            
            return json.dumps({
                "success": False,
                "error": f"Entity '{entity_name}' not found in ResourceDefinitions"
            })
        except Exception as e:
            return json.dumps({"success": False, "error": str(e)})


class CreateResourceLinkTool(BaseTool):
    """Tool for creating resource links."""
    
    name: str = "create_resource_link"
    description: str = "Create or update a resource link entry in the ResourceLinks sheet"
    args_schema: Type[BaseModel] = CreateResourceLinkInput
    state: ToolkitState = Field(default_factory=ToolkitState)
    
    def _run(self, origin_resource: str, reference_path: str, destination_resource: str) -> str:
        try:
            # Use shared workbook instance
            wb = WorkbookManager.get_workbook()
            if wb is None:
                return json.dumps({"success": False, "error": "No workbook is currently open. Use set_current_file or create_new_file first."})
            
            # Validate that both origin and destination resources exist in ResourceDefinitions
            if "ResourceDefinitions" not in wb.sheetnames:
                return json.dumps({"success": False, "error": "ResourceDefinitions sheet not found"})
            
            def_sheet = wb["ResourceDefinitions"]
            existing_entities = set()
            for row in range(3, def_sheet.max_row + 1):
                entity_name = def_sheet.cell(row=row, column=1).value
                if entity_name:
                    existing_entities.add(entity_name)
            
            # Check if origin and destination exist
            errors = []
            if origin_resource not in existing_entities:
                errors.append(f"origin_resource '{origin_resource}' not found in ResourceDefinitions")
            if destination_resource not in existing_entities:
                errors.append(f"destination_resource '{destination_resource}' not found in ResourceDefinitions")
            
            if errors:
                error_msg = "Resource link validation failed:\n" + "\n".join(errors)
                error_msg += f"\n\nExisting resource definitions: {sorted(existing_entities)}"
                return json.dumps({"success": False, "error": error_msg})
            
            if "ResourceLinks" not in wb.sheetnames:
                return json.dumps({"success": False, "error": "ResourceLinks sheet not found"})
            
            sheet = wb["ResourceLinks"]
            existing_row = None
            last_data_row = 2  # Start after header rows
            
            # Find existing link and track last row with actual data
            for row in range(3, sheet.max_row + 1):
                cell_value = sheet.cell(row=row, column=1).value
                if (cell_value == origin_resource and
                    sheet.cell(row=row, column=2).value == reference_path and
                    sheet.cell(row=row, column=3).value == destination_resource):
                    existing_row = row
                    break
                if cell_value is not None and str(cell_value).strip():
                    last_data_row = row
            
            target_row = existing_row if existing_row else last_data_row + 1
            action = "Already exists" if existing_row else "Added"
            
            sheet.cell(row=target_row, column=1, value=origin_resource)
            sheet.cell(row=target_row, column=2, value=reference_path)
            sheet.cell(row=target_row, column=3, value=destination_resource)
            
            return json.dumps({"success": True, "action": action, "row": target_row})
        except Exception as e:
            return json.dumps({"success": False, "error": str(e)})


class AddPatientDataTool(BaseTool):
    """Tool for adding patient data columns."""
    
    name: str = "add_patient_data"
    description: str = "Add a new data column to the PatientData sheet for a specific entity"
    args_schema: Type[BaseModel] = AddPatientDataInput
    state: ToolkitState = Field(default_factory=ToolkitState)
    
    def _run(self, entity_name: str, json_path: str, data_type: str, data_element: str,
             value_set: Optional[str] = None, recommended_profile: Optional[str] = None) -> str:
        try:
            # Use shared workbook instance
            wb = WorkbookManager.get_workbook()
            if wb is None:
                return json.dumps({"success": False, "error": "No workbook is currently open. Use set_current_file or create_new_file first."})
            
            if "PatientData" not in wb.sheetnames:
                return json.dumps({"success": False, "error": "PatientData sheet not found"})
            
            sheet = wb["PatientData"]
            next_col = sheet.max_column + 1
            
            sheet.cell(row=1, column=next_col, value=entity_name)
            sheet.cell(row=2, column=next_col, value=json_path)
            sheet.cell(row=3, column=next_col, value=data_type)
            sheet.cell(row=4, column=next_col, value=value_set if value_set else "")
            sheet.cell(row=5, column=next_col, value=recommended_profile if recommended_profile else "")
            sheet.cell(row=6, column=next_col, value=data_element)
            
            return json.dumps({"success": True, "column": next_col, "entity_name": entity_name})
        except Exception as e:
            return json.dumps({"success": False, "error": str(e)})


class CopyReferenceColumnsTool(BaseTool):
    """Tool for copying reference columns."""
    
    name: str = "copy_reference_columns"
    description: str = "Copy matching columns from ReferenceSheet to PatientData for a specific entity"
    args_schema: Type[BaseModel] = CopyReferenceColumnsInput
    state: ToolkitState = Field(default_factory=ToolkitState)
    
    def _run(self, entity_name: str, value_type: Optional[str] = None) -> str:
        try:
            # Use shared workbook instance
            wb = WorkbookManager.get_workbook()
            if wb is None:
                return json.dumps({"success": False, "error": "No workbook is currently open. Use set_current_file or create_new_file first."})
            
            # Get entity info
            def_sheet = wb["ResourceDefinitions"]
            entity_resource_type = None
            entity_profiles = []
            
            for row in range(3, def_sheet.max_row + 1):
                if def_sheet.cell(row=row, column=1).value == entity_name:
                    entity_resource_type = str(def_sheet.cell(row=row, column=2).value)
                    profiles_str = def_sheet.cell(row=row, column=3).value
                    if profiles_str:
                        entity_profiles = [p.strip() for p in str(profiles_str).split(',')]
                    break
            
            if not entity_resource_type:
                return json.dumps({"success": False, "error": f"Entity '{entity_name}' not found"})
            
            ref_sheet = wb["ReferenceSheet"]
            patient_sheet = wb["PatientData"]
            
            matching_columns = []
            for col in range(1, ref_sheet.max_column + 1):
                json_path = str(ref_sheet.cell(row=2, column=col).value or "")
                if not json_path.startswith(entity_resource_type + "."):
                    continue
                # Only reject if there's a .value path with a mismatched value_type
                # Allow non-value paths regardless of value_type
                if value_type and ".value" in json_path and f".value{value_type}" not in json_path:
                    continue
                
                # Check if the recommended profile (row 5) is in entity_profiles
                recommended_profile = ref_sheet.cell(row=5, column=col).value
                if recommended_profile:
                    recommended_profile = str(recommended_profile).strip()
                    if entity_profiles and recommended_profile not in entity_profiles:
                        continue
                
                matching_columns.append(col)
            
            copied_columns = []
            
            # Find the first truly empty column starting from column C (index 3)
            # by scanning for empty values in rows 1-6
            next_patient_col = 3  # Start from column C
            # Scan up to a reasonable maximum (e.g., 10000 columns)
            for col in range(3, 10000):
                is_empty = True
                for row in range(1, 7):
                    cell_value = patient_sheet.cell(row=row, column=col).value
                    if cell_value is not None and str(cell_value).strip():
                        is_empty = False
                        break
                if is_empty:
                    next_patient_col = col
                    break
            
            for ref_col in matching_columns:
                for row in range(1, 7):
                    ref_cell = ref_sheet.cell(row=row, column=ref_col)
                    patient_cell = patient_sheet.cell(row=row, column=next_patient_col)
                    
                    # Set value (entity name for row 1, otherwise copy from reference)
                    if row == 1:
                        patient_cell.value = entity_name
                    else:
                        patient_cell.value = ref_cell.value
                    
                    # Copy cell style including colors
                    if ref_cell.has_style:
                        patient_cell.font = ref_cell.font.copy()
                        patient_cell.border = ref_cell.border.copy()
                        patient_cell.fill = ref_cell.fill.copy()
                        patient_cell.number_format = ref_cell.number_format
                        patient_cell.protection = ref_cell.protection.copy()
                        patient_cell.alignment = ref_cell.alignment.copy()
                
                copied_columns.append({
                    "column": next_patient_col,
                    "json_path": ref_sheet.cell(row=2, column=ref_col).value
                })
                next_patient_col += 1
            
            return json.dumps({"success": True, "columns_copied": len(copied_columns), "copied_columns": copied_columns})
        except Exception as e:
            return json.dumps({"success": False, "error": str(e)})


class DeleteReferenceColumnsTool(BaseTool):
    """Tool for deleting reference columns."""
    
    name: str = "delete_reference_columns"
    description: str = "Delete all columns in PatientData for a specific entity"
    args_schema: Type[BaseModel] = DeleteReferenceColumnsInput
    state: ToolkitState = Field(default_factory=ToolkitState)
    
    def _run(self, entity_name: str) -> str:
        try:
            # Use shared workbook instance
            wb = WorkbookManager.get_workbook()
            if wb is None:
                return json.dumps({"success": False, "error": "No workbook is currently open. Use set_current_file or create_new_file first."})
            
            patient_sheet = wb["PatientData"]
            
            # Find all columns for this entity
            cols_to_delete = []
            for col in range(3, patient_sheet.max_column + 1):
                if patient_sheet.cell(row=1, column=col).value == entity_name:
                    cols_to_delete.append(col)
            
            # Delete columns in reverse order to maintain indices
            for col in reversed(cols_to_delete):
                patient_sheet.delete_cols(col, 1)
            
            return json.dumps({"success": True, "deleted_columns": len(cols_to_delete), "entity_name": entity_name})
        except Exception as e:
            return json.dumps({"success": False, "error": str(e)})


class CheckResourceDefinitionExistsTool(BaseTool):
    """Tool for checking if a resource definition exists."""
    
    name: str = "check_resource_definition_exists"
    description: str = "Check if a resource definition exists in both ResourceDefinitions and PatientData sheets"
    args_schema: Type[BaseModel] = CheckResourceDefinitionExistsInput
    state: ToolkitState = Field(default_factory=ToolkitState)
    
    def _run(self, entity_name: str) -> str:
        try:
            # Use shared workbook instance
            wb = WorkbookManager.get_workbook()
            if wb is None:
                return json.dumps({"success": False, "error": "No workbook is currently open. Use set_current_file or create_new_file first."})
            
            # Check ResourceDefinitions sheet
            exists_in_definitions = False
            resource_type = None
            profiles = None
            
            if "ResourceDefinitions" in wb.sheetnames:
                def_sheet = wb["ResourceDefinitions"]
                for row in range(3, def_sheet.max_row + 1):
                    if def_sheet.cell(row=row, column=1).value == entity_name:
                        exists_in_definitions = True
                        resource_type = def_sheet.cell(row=row, column=2).value
                        profiles = def_sheet.cell(row=row, column=3).value
                        break
            
            # Check PatientData sheet
            exists_in_patient_data = False
            column_count = 0
            
            if "PatientData" in wb.sheetnames:
                patient_sheet = wb["PatientData"]
                for col in range(3, patient_sheet.max_column + 1):
                    if patient_sheet.cell(row=1, column=col).value == entity_name:
                        exists_in_patient_data = True
                        column_count += 1
            
            # Determine overall existence
            exists = exists_in_definitions and exists_in_patient_data
            
            return json.dumps({
                "success": True,
                "exists": exists,
                "entity_name": entity_name,
                "exists_in_definitions": exists_in_definitions,
                "exists_in_patient_data": exists_in_patient_data,
                "resource_type": resource_type,
                "profiles": profiles if profiles else "",
                "patient_data_column_count": column_count
            })
        except Exception as e:
            return json.dumps({"success": False, "error": str(e)})


class SetPatientDataValueTool(BaseTool):
    """Tool for setting patient data values."""
    
    name: str = "set_patient_data_value"
    description: str = "Set a specific patient data value in the PatientData sheet"
    args_schema: Type[BaseModel] = SetPatientDataValueInput
    state: ToolkitState = Field(default_factory=ToolkitState)
    
    def _run(self, entity_name: str, data_element: str, row_index: int, value: str) -> str:
        try:
            # Use shared workbook instance
            wb = WorkbookManager.get_workbook()
            if wb is None:
                return json.dumps({"success": False, "error": "No workbook is currently open. Use set_current_file or create_new_file first."})
            
            sheet = wb["PatientData"]
            
            target_column = None
            for col in range(1, sheet.max_column + 1):
                if (sheet.cell(row=1, column=col).value == entity_name and
                    sheet.cell(row=6, column=col).value == data_element):
                    target_column = col
                    break
            
            if target_column is None:
                return json.dumps({"success": False, "error": f"Column not found for {entity_name}.{data_element}"})
            
            excel_row = row_index + 7
            sheet.cell(row=excel_row, column=target_column, value=value)
            
            return json.dumps({"success": True, "column": target_column, "row_index": row_index, "excel_row": excel_row})
        except Exception as e:
            return json.dumps({"success": False, "error": str(e)})


class GenerateFhirBundlesTool(BaseTool):
    """Tool for generating FHIR bundles."""
    
    name: str = "generate_fhir_bundles"
    description: str = "Generate FHIR bundle JSON files from the Excel file using the fhir-sheets CLI tool"
    args_schema: Type[BaseModel] = GenerateFhirBundlesInput
    state: ToolkitState = Field(default_factory=ToolkitState)
    
    def _run(self, output_dir: str = "./output", config_path: Optional[str] = None) -> str:
        try:
            from fhir_sheets.cli.main import main as fhir_sheets_main
            from fhir_sheets.core.config.FhirSheetsConfiguration import FhirSheetsConfiguration
            
            # Use the file path from WorkbookManager
            xlsx_path = WorkbookManager.get_file_path()
            if xlsx_path is None:
                return json.dumps({"success": False, "error": "No workbook is currently open. Use set_current_file or create_new_file first."})
            
            output_path = Path(output_dir)
            output_path.mkdir(parents=True, exist_ok=True)
            
            config_file = config_path or self.state.config_path
            if config_file:
                with open(config_file, 'r') as f:
                    config_dict = json.load(f)
                config = FhirSheetsConfiguration(config_dict)
            else:
                config = FhirSheetsConfiguration({})
            
            fhir_sheets_main(xlsx_path, output_dir, 'bundle', config)
            files_created = len(list(output_path.glob('*.json')))
            
            return json.dumps({"success": True, "files_created": files_created, "output_dir": output_dir})
        except Exception as e:
            return json.dumps({"success": False, "error": str(e)})


class GetMetadataTool(BaseTool):
    """Tool for getting metadata."""
    
    name: str = "get_metadata"
    description: str = "Extract metadata from a FHIR cohort Excel file"
    args_schema: Type[BaseModel] = GetMetadataInput
    state: ToolkitState = Field(default_factory=ToolkitState)
    
    def _run(self) -> str:
        try:
            # Use shared workbook instance
            wb = WorkbookManager.get_workbook()
            if wb is None:
                return json.dumps({"error": "No workbook is currently open. Use set_current_file or create_new_file first."})
            
            xlsx_path = WorkbookManager.get_file_path()
            sheets = wb.sheetnames
            
            resource_definitions = []
            if "ResourceDefinitions" in sheets:
                sheet = wb["ResourceDefinitions"]
                for row in range(3, sheet.max_row + 1):
                    entity_name = sheet.cell(row=row, column=1).value
                    if entity_name:
                        resource_definitions.append({
                            "entity_name": entity_name,
                            "resource_type": sheet.cell(row=row, column=2).value
                        })
            
            return json.dumps({
                "xlsx_path": xlsx_path,
                "sheets": sheets,
                "resource_definitions": resource_definitions,
                "resource_definitions_count": len(resource_definitions)
            })
        except Exception as e:
            return json.dumps({"error": str(e)})


# Toolkit Class
class FhirSheetsXlsxToolkit(BaseToolkit):
    """Toolkit for FHIRSheets XLSX operations."""
    
    state: ToolkitState = Field(default=ToolkitState(), description="Shared state for the toolkit")
    
    def __init__(self, config_path: Optional[str] = None, working_dir: Optional[str] = None):
        """Initialize the toolkit with shared state."""
        super().__init__()
        self.state = ToolkitState()
        if config_path:
            self.state.config_path = config_path
        if working_dir:
            self.state.working_dir = Path(working_dir)
    
    def get_tools(self) -> List[BaseTool]:
        """Get all FHIRSheets XLSX tools with shared state."""
        return [
            CreateNewFileTool(state=self.state),
            SetCurrentFileTool(state=self.state),
            GetCurrentFileTool(state=self.state),
            CreateResourceDefinitionTool(state=self.state),
            GetResourceDefinitionTool(state=self.state),
            CreateResourceLinkTool(state=self.state),
            AddPatientDataTool(state=self.state),
            CopyReferenceColumnsTool(state=self.state),
            DeleteReferenceColumnsTool(state=self.state),
            CheckResourceDefinitionExistsTool(state=self.state),
            SetPatientDataValueTool(state=self.state),
            GenerateFhirBundlesTool(state=self.state),
            GetMetadataTool(state=self.state),
        ]
