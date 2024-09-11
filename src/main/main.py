import openpyxl

# Function to process each row and convert to a new format
def process_row(headers, row):
    # Create a dictionary where headers are keys and row values are values
    return {header: value for header, value in zip(headers, row)}

# Read the xlsx file with custom row and column ranges
def read_and_convert_xlsx(file_path):
    # Load the workbook and select the active sheet
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active

    # Extract headers from row 5, starting from column 2
    headers = [cell.value for cell in sheet[5][1:]]

    converted_rows = []

    # Extract data from rows 6-15 (inclusive), starting from column 2
    for row in sheet.iter_rows(min_row=6, max_row=15, min_col=2, values_only=True):
        converted = process_row(headers, row)
        converted_rows.append(converted)
    return converted_rows

# Replace placeholders in the template
def replace_placeholders(template_text, row_data):
    for key, value in row_data.items():
        placeholder = f"${{{key}}}"
        template_text = template_text.replace(placeholder, str(value))
    return template_text

# Read the JSON template
def read_json_template(file_path):
    with open(file_path, 'r') as file:
        return file.read()

# Read input file and convert to local format
input_xlsx_file_path = '.\\src\\resources\\Synthetic_Input.xlsx'
patient_json_template_path = '.\\src\\resources\\json_templates\\Patient.json'
converted_data = read_and_convert_xlsx(input_xlsx_file_path)

# Read the JSON template
patient_template_text = read_json_template(patient_json_template_path)


# Iterate over the rows and replace placeholders
for row_data in converted_data:
    filled_template = replace_placeholders(patient_template_text, row_data)
    print(filled_template)  # Or save it to a file if needed