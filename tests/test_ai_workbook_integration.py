"""
Tests for the integrated AI workbook management workflow
"""

import json
import pytest
import tempfile
import os
from pathlib import Path
from openpyxl import Workbook
from langgraph_dev.tools.fhirsheets_xlsx_tool import (
    WorkbookManager,
    FhirSheetsXlsxToolkit
)


@pytest.fixture
def temp_workbook():
    """Create a temporary workbook for testing."""
    # Create a temporary file
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    
    # Create a simple workbook with required sheets
    wb = Workbook()
    
    # Create ResourceDefinitions sheet
    ws = wb.active
    ws.title = "ResourceDefinitions"
    ws['A1'] = "Entity Name"
    ws['B1'] = "Resource Type"
    ws['C1'] = "Profiles"
    
    # Create PatientData sheet
    wb.create_sheet("PatientData")
    
    wb.save(path)
    wb.close()
    
    yield path
    
    # Cleanup
    if WorkbookManager.is_open():
        WorkbookManager.close_workbook(save=False)
    if os.path.exists(path):
        os.remove(path)


def test_set_current_file_opens_workbook(temp_workbook):
    """Test that set_current_file opens the workbook in WorkbookManager."""
    # Ensure no workbook is open
    if WorkbookManager.is_open():
        WorkbookManager.close_workbook(save=False)
    
    toolkit = FhirSheetsXlsxToolkit()
    tools = toolkit.get_tools()
    set_file_tool = next(t for t in tools if t.name == "set_current_file")
    
    # Set current file
    result = set_file_tool.invoke({"file_path": temp_workbook})
    result_dict = json.loads(result)
    
    assert result_dict["success"] is True
    assert WorkbookManager.is_open()
    assert WorkbookManager.get_file_path() == temp_workbook
    
    # Cleanup
    WorkbookManager.close_workbook(save=False)


def test_create_resource_definition_uses_shared_workbook(temp_workbook):
    """Test that create_resource_definition uses the shared workbook."""
    # Ensure no workbook is open
    if WorkbookManager.is_open():
        WorkbookManager.close_workbook(save=False)
    
    toolkit = FhirSheetsXlsxToolkit()
    tools = toolkit.get_tools()
    set_file_tool = next(t for t in tools if t.name == "set_current_file")
    create_def_tool = next(t for t in tools if t.name == "create_resource_definition")
    
    # Set current file (opens workbook)
    set_file_tool.invoke({"file_path": temp_workbook})
    
    # Create resource definition (uses shared workbook)
    result = create_def_tool.invoke({
        "entity_name": "TestPatient",
        "resource_type": "Patient"
    })
    result_dict = json.loads(result)
    
    assert result_dict["success"] is True
    assert result_dict["entity_name"] == "TestPatient"
    assert result_dict["resource_type"] == "Patient"
    
    # Verify workbook is still open
    assert WorkbookManager.is_open()
    
    # Cleanup
    WorkbookManager.close_workbook(save=False)


def test_operations_without_open_workbook_fail():
    """Test that operations fail when no workbook is open."""
    # Ensure no workbook is open
    if WorkbookManager.is_open():
        WorkbookManager.close_workbook(save=False)
    
    toolkit = FhirSheetsXlsxToolkit()
    tools = toolkit.get_tools()
    create_def_tool = next(t for t in tools if t.name == "create_resource_definition")
    
    # Try to create resource definition without opening workbook
    result = create_def_tool.invoke({
        "entity_name": "TestPatient",
        "resource_type": "Patient"
    })
    result_dict = json.loads(result)
    
    assert result_dict["success"] is False
    assert "No workbook is currently open" in result_dict["error"]


def test_workbook_auto_switches_on_set_current_file(temp_workbook):
    """Test that setting a different file closes the previous one."""
    # Create a second temp workbook
    fd, path2 = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb = Workbook()
    wb.save(path2)
    wb.close()
    
    try:
        # Ensure no workbook is open
        if WorkbookManager.is_open():
            WorkbookManager.close_workbook(save=False)
        
        toolkit = FhirSheetsXlsxToolkit()
        tools = toolkit.get_tools()
        set_file_tool = next(t for t in tools if t.name == "set_current_file")
        
        # Open first file
        set_file_tool.invoke({"file_path": temp_workbook})
        assert WorkbookManager.get_file_path() == temp_workbook
        
        # Open second file (should close first and open second)
        set_file_tool.invoke({"file_path": path2})
        assert WorkbookManager.get_file_path() == path2
        
        # Cleanup
        WorkbookManager.close_workbook(save=False)
    finally:
        if os.path.exists(path2):
            os.remove(path2)