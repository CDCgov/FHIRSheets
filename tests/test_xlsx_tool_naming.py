"""Unit tests for XLSX tool naming standards."""

import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from src.langgraph_dev.tools.fhirsheets_xlsx_tool import (
    CreateNewFileTool,
    CreateResourceDefinitionTool,
    CreateResourceLinkTool,
    ToolkitState,
    WorkbookManager,
    _to_snake_case,
    _to_pascal_case,
)


class TestNamingUtilities(unittest.TestCase):
    """Test naming utility functions."""
    
    def test_to_snake_case_spaces(self):
        """Test conversion of spaces to snake_case."""
        self.assertEqual(_to_snake_case("My Test File"), "my_test_file")
    
    def test_to_snake_case_pascal(self):
        """Test conversion of PascalCase to snake_case."""
        self.assertEqual(_to_snake_case("MyTestFile"), "my_test_file")
    
    def test_to_snake_case_camel(self):
        """Test conversion of camelCase to snake_case."""
        self.assertEqual(_to_snake_case("myTestFile"), "my_test_file")
    
    def test_to_snake_case_hyphens(self):
        """Test conversion of hyphens to snake_case."""
        self.assertEqual(_to_snake_case("my-test-file"), "my_test_file")
    
    def test_to_snake_case_mixed(self):
        """Test conversion of mixed formats to snake_case."""
        self.assertEqual(_to_snake_case("My-Test File"), "my_test_file")
    
    def test_to_pascal_case_snake(self):
        """Test conversion of snake_case to PascalCase."""
        self.assertEqual(_to_pascal_case("my_patient"), "MyPatient")
    
    def test_to_pascal_case_spaces(self):
        """Test conversion of spaces to PascalCase."""
        self.assertEqual(_to_pascal_case("my patient"), "MyPatient")
    
    def test_to_pascal_case_hyphens(self):
        """Test conversion of hyphens to PascalCase."""
        self.assertEqual(_to_pascal_case("my-patient"), "MyPatient")
    
    def test_to_pascal_case_already_pascal(self):
        """Test that PascalCase names remain unchanged."""
        output = _to_pascal_case("MyPatient")
        self.assertEqual(output, "MyPatient")


class TestCreateResourceDefinitionTool(unittest.TestCase):
    """Test CreateResourceDefinitionTool with PascalCase conversion."""
    
    def setUp(self):
        """Set up test fixtures."""
        self.temp_dir = tempfile.mkdtemp()
        self.state = ToolkitState(working_dir=Path(self.temp_dir))
        self.tool = CreateResourceDefinitionTool(state=self.state)
        
        # Create a test Excel file
        self.test_file = Path(self.temp_dir) / "test.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.create_sheet("ResourceDefinitions")
        ws.cell(row=1, column=1, value="Entity Name")
        ws.cell(row=1, column=2, value="ResourceType")
        ws.cell(row=1, column=3, value="Profile(s)")
        wb.save(self.test_file)
        wb.close()
    
    def tearDown(self):
        """Clean up test fixtures."""
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)
    
    def test_snake_case_converted_to_pascal(self):
        """Test that snake_case names are converted to PascalCase."""
        WorkbookManager.open_workbook(str(self.test_file))
        
        result = self.tool._run(
            entity_name="my_patient",
            resource_type="Patient"
        )
        
        WorkbookManager.close_workbook(save=False)
        result_dict = json.loads(result)
        self.assertTrue(result_dict.get("success"))
        self.assertEqual(result_dict.get("entity_name"), "MyPatient")
    
    def test_spaces_converted_to_pascal(self):
        """Test that names with spaces are converted to PascalCase."""
        WorkbookManager.open_workbook(str(self.test_file))
        
        result = self.tool._run(
            entity_name="my patient",
            resource_type="Patient"
        )
        
        WorkbookManager.close_workbook(save=False)
        result_dict = json.loads(result)
        self.assertTrue(result_dict.get("success"))
        self.assertEqual(result_dict.get("entity_name"), "MyPatient")
    
    def test_pascal_case_unchanged(self):
        """Test that PascalCase names remain unchanged."""
        WorkbookManager.open_workbook(str(self.test_file))
        
        result = self.tool._run(
            entity_name="MyPatient",
            resource_type="Patient"
        )
        
        WorkbookManager.close_workbook(save=False)
        result_dict = json.loads(result)
        self.assertTrue(result_dict.get("success"))
        self.assertEqual(result_dict.get("entity_name"), "MyPatient")


class TestCreateResourceLinkTool(unittest.TestCase):
    """Test CreateResourceLinkTool with resource validation."""
    
    def setUp(self):
        """Set up test fixtures."""
        self.temp_dir = tempfile.mkdtemp()
        self.state = ToolkitState(working_dir=Path(self.temp_dir))
        self.tool = CreateResourceLinkTool(state=self.state)
        
        # Create a test Excel file with resource definitions
        self.test_file = Path(self.temp_dir) / "test.xlsx"
        wb = openpyxl.Workbook()
        
        # ResourceDefinitions sheet
        def_ws = wb.create_sheet("ResourceDefinitions")
        def_ws.cell(row=1, column=1, value="Entity Name")
        def_ws.cell(row=1, column=2, value="ResourceType")
        def_ws.cell(row=3, column=1, value="MyPatient")
        def_ws.cell(row=3, column=2, value="Patient")
        def_ws.cell(row=4, column=1, value="MyEncounter")
        def_ws.cell(row=4, column=2, value="Encounter")
        
        # ResourceLinks sheet
        link_ws = wb.create_sheet("ResourceLinks")
        link_ws.cell(row=1, column=1, value="OriginResource")
        link_ws.cell(row=1, column=2, value="ReferencePath")
        link_ws.cell(row=1, column=3, value="DestinationResource")
        
        wb.save(self.test_file)
        wb.close()
    
    def tearDown(self):
        """Clean up test fixtures."""
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)
    
    def test_valid_resource_link(self):
        """Test that links between existing resources are accepted."""
        WorkbookManager.open_workbook(str(self.test_file))
        
        result = self.tool._run(
            origin_resource="MyEncounter",
            reference_path="subject",
            destination_resource="MyPatient"
        )
        
        WorkbookManager.close_workbook(save=False)
        result_dict = json.loads(result)
        self.assertTrue(result_dict.get("success"))
    
    def test_invalid_origin_resource(self):
        """Test that links with non-existent origin are rejected."""
        WorkbookManager.open_workbook(str(self.test_file))
        
        result = self.tool._run(
            origin_resource="NonExistent",
            reference_path="subject",
            destination_resource="MyPatient"
        )
        
        WorkbookManager.close_workbook(save=False)
        result_dict = json.loads(result)
        self.assertFalse(result_dict.get("success"))
        self.assertIn("NonExistent", result_dict.get("error"))
        self.assertIn("Existing resource definitions", result_dict.get("error"))
    
    def test_invalid_destination_resource(self):
        """Test that links with non-existent destination are rejected."""
        WorkbookManager.open_workbook(str(self.test_file))
        
        result = self.tool._run(
            origin_resource="MyEncounter",
            reference_path="subject",
            destination_resource="NonExistent"
        )
        
        WorkbookManager.close_workbook(save=False)
        result_dict = json.loads(result)
        self.assertFalse(result_dict.get("success"))
        self.assertIn("NonExistent", result_dict.get("error"))
        self.assertIn("Existing resource definitions", result_dict.get("error"))


if __name__ == "__main__":
    unittest.main()